#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Відкриваємо існуючу презентацію
prs = Presentation('/home/user/demo2/n8n-presentation.pptx')

# Slide: API
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(239, 246, 255)

title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
tf = title.text_frame
tf.text = "API - Application Programming Interface"
tf.paragraphs[0].font.size = Pt(38)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(31, 41, 55)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

subtitle = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(0.5))
stf = subtitle.text_frame
stf.text = "Аналогія: Офіціант у ресторані"
stf.paragraphs[0].font.size = Pt(24)
stf.paragraphs[0].font.bold = True
stf.paragraphs[0].font.color.rgb = RGBColor(37, 99, 235)
stf.paragraphs[0].alignment = PP_ALIGN.CENTER

# 4 boxes
boxes = [
    {"icon": "👤", "title": "Ви (n8n)", "text": "Не йдете на кухню (в базу даних Google)", "left": 0.5},
    {"icon": "👨‍💼", "title": "Офіціант (API)", "text": "Ви кличете офіціанта", "left": 2.75},
    {"icon": "📋", "title": "Чітке замовлення", "text": "Даєте замовлення 'Отримати мої події'", "left": 5},
    {"icon": "🍽️", "title": "Готова страва", "text": "Офіціант приносить готові дані", "left": 7.25}
]

for box in boxes:
    # Box
    box_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(box["left"]), Inches(2.5), Inches(2), Inches(3))
    box_shape.fill.solid()
    box_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box_shape.line.color.rgb = RGBColor(59, 130, 246)
    box_shape.line.width = Pt(2)

    # Icon
    icon_box = slide.shapes.add_textbox(Inches(box["left"] + 0.1), Inches(2.7), Inches(1.8), Inches(0.5))
    tf = icon_box.text_frame
    tf.text = box["icon"]
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Title
    title_box = slide.shapes.add_textbox(Inches(box["left"] + 0.1), Inches(3.3), Inches(1.8), Inches(0.4))
    tf = title_box.text_frame
    tf.text = box["title"]
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(31, 41, 55)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Text
    text_box = slide.shapes.add_textbox(Inches(box["left"] + 0.1), Inches(3.9), Inches(1.8), Inches(1.3))
    tf = text_box.text_frame
    tf.text = box["text"]
    tf.word_wrap = True
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = RGBColor(75, 85, 99)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Slide: JSON
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(239, 246, 255)

title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
tf = title.text_frame
tf.text = "JSON - JavaScript Object Notation"
tf.paragraphs[0].font.size = Pt(38)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(31, 41, 55)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

subtitle = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(0.5))
stf = subtitle.text_frame
stf.text = "Аналогія: Структурований список покупок"
stf.paragraphs[0].font.size = Pt(20)
stf.paragraphs[0].font.color.rgb = RGBColor(107, 114, 128)
stf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Left side - explanation
left_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(4), Inches(3))
tf = left_box.text_frame
tf.text = "Це мова, якою ви 'пишете замовлення' для офіціанта (API).\n\nЧітка структура:\n\"ключ\": \"значення\""
tf.paragraphs[0].font.size = Pt(18)
tf.paragraphs[0].font.color.rgb = RGBColor(75, 85, 99)

# Right side - JSON example
json_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5), Inches(2.2), Inches(4.5), Inches(3.5))
json_box.fill.solid()
json_box.fill.fore_color.rgb = RGBColor(30, 30, 46)
json_box.line.color.rgb = RGBColor(100, 100, 120)

json_text = slide.shapes.add_textbox(Inches(5.3), Inches(2.5), Inches(3.9), Inches(3))
tf = json_text.text_frame
json_code = """[
  {
    "категорія": "Молочні",
    "продукт": "Молоко",
    "кількість": 1
  },
  ...
]"""
tf.text = json_code
tf.paragraphs[0].font.name = 'Courier New'
tf.paragraphs[0].font.size = Pt(14)
tf.paragraphs[0].font.color.rgb = RGBColor(166, 227, 161)

# Slide: JSON вкладені дані
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(239, 246, 255)

title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
tf = title.text_frame
tf.text = "JSON: Вкладені дані"
tf.paragraphs[0].font.size = Pt(40)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(31, 41, 55)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Explanation
expl_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4), Inches(2))
tf = expl_box.text_frame
tf.text = "Аналогія: Рецепт\n(Об'єкт + Список)\n\n{ } = Об'єкт (один рецепт)\n[ ] = Список (інгредієнтів)"
tf.paragraphs[0].font.size = Pt(18)
tf.paragraphs[0].font.color.rgb = RGBColor(75, 85, 99)

# JSON example
json_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5), Inches(1.5), Inches(4.5), Inches(4.5))
json_box.fill.solid()
json_box.fill.fore_color.rgb = RGBColor(30, 30, 46)

json_text = slide.shapes.add_textbox(Inches(5.3), Inches(1.8), Inches(3.9), Inches(4))
tf = json_text.text_frame
json_code = """{
  "назва": "Омлет",
  "інгредієнти": [
    {
      "назва": "Яйця",
      "кількість": 2
    },
    {
      "назва": "Молоко",
      "кількість": 50
    }
  ]
}"""
tf.text = json_code
tf.paragraphs[0].font.name = 'Courier New'
tf.paragraphs[0].font.size = Pt(13)
tf.paragraphs[0].font.color.rgb = RGBColor(166, 227, 161)

# Slide: Аналогія з поштою
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(239, 246, 255)

title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
tf = title.text_frame
tf.text = "Аналогія: Пошта"
tf.paragraphs[0].font.size = Pt(44)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(31, 41, 55)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Left box - One element
left_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2), Inches(4.5), Inches(3))
left_box.fill.solid()
left_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
left_box.line.color.rgb = RGBColor(59, 130, 246)
left_box.line.width = Pt(3)

title_left = slide.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(4.1), Inches(0.5))
tf = title_left.text_frame
tf.text = "📧 Один Елемент"
tf.paragraphs[0].font.size = Pt(24)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(59, 130, 246)

text_left = slide.shapes.add_textbox(Inches(0.7), Inches(2.9), Inches(4.1), Inches(1.9))
tf = text_left.text_frame
content = "✓ Gmail знайшла 1 лист\n\n→ Передає 1 конверт\n\n→ Наступна нода спрацює 1 РАЗ"
tf.text = content
tf.paragraphs[0].font.size = Pt(16)
tf.paragraphs[0].font.color.rgb = RGBColor(55, 65, 81)

# Right box - List
right_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.5), Inches(2), Inches(4.5), Inches(3))
right_box.fill.solid()
right_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
right_box.line.color.rgb = RGBColor(249, 115, 22)
right_box.line.width = Pt(3)

title_right = slide.shapes.add_textbox(Inches(5.7), Inches(2.2), Inches(4.1), Inches(0.5))
tf = title_right.text_frame
tf.text = "📚 Список Елементів"
tf.paragraphs[0].font.size = Pt(24)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(249, 115, 22)

text_right = slide.shapes.add_textbox(Inches(5.7), Inches(2.9), Inches(4.1), Inches(1.9))
tf = text_right.text_frame
content = "✓ Google Sheets знайшла 10 рядків\n\n→ Передає стос з 10 конвертів\n\n→ Наступна нода спрацює 10 РАЗІВ"
tf.text = content
tf.paragraphs[0].font.size = Pt(16)
tf.paragraphs[0].font.color.rgb = RGBColor(55, 65, 81)

# Conclusion
conclusion = slide.shapes.add_textbox(Inches(1.5), Inches(5.8), Inches(7), Inches(0.8))
tf = conclusion.text_frame
tf.text = "💡 Ключове правило: Кількість елементів = Кількість виконань"
tf.paragraphs[0].font.size = Pt(20)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(126, 34, 206)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Slide: Webhook
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(239, 246, 255)

title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
tf = title.text_frame
tf.text = "WEBHOOK"
tf.paragraphs[0].font.size = Pt(44)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(31, 41, 55)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

subtitle = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(0.5))
stf = subtitle.text_frame
stf.text = "Аналогія: Дверний дзвінок"
stf.paragraphs[0].font.size = Pt(24)
stf.paragraphs[0].font.color.rgb = RGBColor(107, 114, 128)
stf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Comparison boxes
left_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2.5), Inches(4), Inches(3))
left_box.fill.solid()
left_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
left_box.line.color.rgb = RGBColor(249, 115, 22)
left_box.line.width = Pt(3)

title_left = slide.shapes.add_textbox(Inches(1.2), Inches(2.7), Inches(3.6), Inches(0.5))
tf = title_left.text_frame
tf.text = "📞 Звичайний API"
tf.paragraphs[0].font.size = Pt(20)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(249, 115, 22)

text_left = slide.shapes.add_textbox(Inches(1.2), Inches(3.3), Inches(3.6), Inches(1.9))
tf = text_left.text_frame
content = "Ви кожні 5 хвилин дзвоните в піцерію:\n\n'Моя піца готова? А зараз?'\n\n⚠️ Постійні перевірки, марнування часу"
tf.text = content
tf.paragraphs[0].font.size = Pt(14)
tf.paragraphs[0].font.color.rgb = RGBColor(75, 85, 99)

right_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.5), Inches(2.5), Inches(4), Inches(3))
right_box.fill.solid()
right_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
right_box.line.color.rgb = RGBColor(34, 197, 94)
right_box.line.width = Pt(3)

title_right = slide.shapes.add_textbox(Inches(5.7), Inches(2.7), Inches(3.6), Inches(0.5))
tf = title_right.text_frame
tf.text = "🔔 Webhook"
tf.paragraphs[0].font.size = Pt(20)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(34, 197, 94)

text_right = slide.shapes.add_textbox(Inches(5.7), Inches(3.3), Inches(3.6), Inches(1.9))
tf = text_right.text_frame
content = "Ви кажете:\n\n'Ось номер мого дзвінка. Подзвоніть, коли буде готово'\n\n✓ Очікування на сповіщення, ефективно"
tf.text = content
tf.paragraphs[0].font.size = Pt(14)
tf.paragraphs[0].font.color.rgb = RGBColor(75, 85, 99)

conclusion = slide.shapes.add_textbox(Inches(1.5), Inches(6), Inches(7), Inches(0.7))
tf = conclusion.text_frame
tf.text = "Webhook — це URL, який чекає на дані"
tf.paragraphs[0].font.size = Pt(20)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(34, 197, 94)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Slide: API Token
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(249, 250, 251)

title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
tf = title.text_frame
tf.text = "🔑 API Токен"
tf.paragraphs[0].font.size = Pt(44)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(31, 41, 55)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Left side
left_text = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(5), Inches(4))
tf = left_text.text_frame

points = [
    "👨‍💼 Офіціант (API) має знати, хто ви",
    "",
    "🆔 Токен (довгий рядок) — це ваш доказ: 'Я — це я'",
    "",
    "🔒 Це ваш секретний пароль!",
    "❗️ Нікому не показуйте ❗️"
]

for i, point in enumerate(points):
    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
    p.text = point
    p.font.size = Pt(18)
    if "секретний" in point or "Нікому" in point:
        p.font.color.rgb = RGBColor(220, 38, 38)
        p.font.bold = True
    else:
        p.font.color.rgb = RGBColor(55, 65, 81)
    p.space_after = Pt(10)

# Right side - Token example
token_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.8), Inches(2.5), Inches(3.7), Inches(3))
token_box.fill.solid()
token_box.fill.fore_color.rgb = RGBColor(251, 191, 36)

crown = slide.shapes.add_textbox(Inches(5.8), Inches(2.7), Inches(3.7), Inches(0.8))
tf = crown.text_frame
tf.text = "👑"
tf.paragraphs[0].font.size = Pt(48)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

token_text_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6), Inches(3.7), Inches(3.3), Inches(1.5))
token_text_box.fill.solid()
token_text_box.fill.fore_color.rgb = RGBColor(255, 255, 255)

token_text = slide.shapes.add_textbox(Inches(6.1), Inches(3.85), Inches(3.1), Inches(1.3))
tf = token_text.text_frame
tf.text = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJzdWIiOiIxMjM..."
tf.word_wrap = True
tf.paragraphs[0].font.name = 'Courier New'
tf.paragraphs[0].font.size = Pt(9)
tf.paragraphs[0].font.color.rgb = RGBColor(75, 85, 99)

# Slide: HTTP Запити
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(249, 250, 251)

title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
tf = title.text_frame
tf.text = "🌐 HTTP Запити"
tf.paragraphs[0].font.size = Pt(44)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(31, 41, 55)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

subtitle = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(0.5))
stf = subtitle.text_frame
stf.text = "Аналогія: Тип вашого замовлення"
stf.paragraphs[0].font.size = Pt(24)
stf.paragraphs[0].font.color.rgb = RGBColor(107, 114, 128)
stf.paragraphs[0].alignment = PP_ALIGN.CENTER

# HTTP methods
methods = [
    {"name": "GET", "icon": "⬇️", "desc": "Отримати\n'Дай мені мої листи'", "color": (34, 197, 94), "left": 1, "top": 2.5},
    {"name": "POST", "icon": "➕", "desc": "Створити\n'Візьми це і створи новий пост'", "color": (59, 130, 246), "left": 5.5, "top": 2.5},
    {"name": "PUT", "icon": "✏️", "desc": "Оновити\n'Зміни назву цієї події'", "color": (249, 115, 22), "left": 1, "top": 4.8},
    {"name": "DELETE", "icon": "🗑️", "desc": "Видалити\n'Видали цей файл'", "color": (239, 68, 68), "left": 5.5, "top": 4.8}
]

for method in methods:
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(method["left"]), Inches(method["top"]), Inches(4), Inches(1.8))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box.line.color.rgb = RGBColor(*method["color"])
    box.line.width = Pt(3)

    # Method name and icon
    header = slide.shapes.add_textbox(Inches(method["left"] + 0.2), Inches(method["top"] + 0.2), Inches(3.6), Inches(0.5))
    tf = header.text_frame
    tf.text = f"{method['icon']} {method['name']}"
    tf.paragraphs[0].font.size = Pt(24)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(*method["color"])

    # Description
    desc = slide.shapes.add_textbox(Inches(method["left"] + 0.2), Inches(method["top"] + 0.8), Inches(3.6), Inches(0.9))
    tf = desc.text_frame
    tf.text = method["desc"]
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.color.rgb = RGBColor(55, 65, 81)

print("Додано слайди про API, JSON, Webhook, Токен, HTTP запити")

# Save
prs.save('/home/user/demo2/n8n-presentation.pptx')
print(f"✓ Оновлено презентацію. Всього слайдів: {len(prs.slides)}")
