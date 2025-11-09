#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import requests
from io import BytesIO

# Відкриваємо існуючу презентацію
prs = Presentation('/home/user/demo2/n8n-presentation.pptx')

# Slide: Базові блоки (Ноди) в n8n
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(249, 250, 251)

title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
tf = title.text_frame
tf.text = "Базові блоки (Ноди) в n8n"
tf.paragraphs[0].font.size = Pt(40)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(31, 41, 55)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Three columns
columns = [
    {
        "title": "🔴 Тригери",
        "color": (239, 68, 68),
        "items": [
            "👆 Manual (Start)",
            "⏰ Schedule (Cron)",
            "🌐 Webhook"
        ],
        "left": 0.8
    },
    {
        "title": "🧠 Логіка",
        "color": (249, 115, 22),
        "items": [
            "❓ IF (Якщо)",
            "🔀 Switch (Вибір)",
            "✏️ Set (Змінити дані)"
        ],
        "left": 3.8
    },
    {
        "title": "✅ Дії",
        "color": (34, 197, 94),
        "items": [
            "📊 Google Sheets",
            "📧 Gmail / Email",
            "🌐 HTTP Request",
            "🤖 OpenAI (AI)"
        ],
        "left": 6.8
    }
]

for col in columns:
    # Column box
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(col["left"]), Inches(2), Inches(2.8), Inches(4))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box.line.color.rgb = RGBColor(*col["color"])
    box.line.width = Pt(3)

    # Title
    title_box = slide.shapes.add_textbox(Inches(col["left"] + 0.2), Inches(2.2), Inches(2.4), Inches(0.6))
    tf = title_box.text_frame
    tf.text = col["title"]
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(*col["color"])
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Items
    items_box = slide.shapes.add_textbox(Inches(col["left"] + 0.3), Inches(3), Inches(2.2), Inches(2.8))
    tf = items_box.text_frame

    for i, item in enumerate(col["items"]):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = item
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(55, 65, 81)
        p.space_after = Pt(15)

# Додатковий слайд з детальною інформацією про популярні ноди
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(239, 246, 255)

title = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.7))
tf = title.text_frame
tf.text = "Популярні ноди n8n"
tf.paragraphs[0].font.size = Pt(40)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(31, 41, 55)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Node descriptions
nodes = [
    {
        "icon": "🌐",
        "name": "HTTP Request",
        "desc": "Виконує запити до будь-якого API. Підтримує GET, POST, PUT, DELETE з аутентифікацією",
        "color": (59, 130, 246),
        "top": 1.5,
        "left": 0.5
    },
    {
        "icon": "🔔",
        "name": "Webhook",
        "desc": "Приймає дані в реальному часі. Створює унікальний URL для кожного workflow",
        "color": (168, 85, 247),
        "top": 1.5,
        "left": 5.5
    },
    {
        "icon": "💻",
        "name": "Code (JavaScript)",
        "desc": "Виконує JavaScript код для обробки даних. Повний доступ до ES6+ та бібліотек",
        "color": (34, 197, 94),
        "top": 3.2,
        "left": 0.5
    },
    {
        "icon": "❓",
        "name": "IF Node",
        "desc": "Розгалужує workflow на основі умов. Підтримує порівняння, regex, AND/OR",
        "color": (251, 191, 36),
        "top": 3.2,
        "left": 5.5
    },
    {
        "icon": "⏰",
        "name": "Schedule Trigger",
        "desc": "Запускає workflow автоматично: щохвилини, щогодини, за cron-розкладом",
        "color": (99, 102, 241),
        "top": 4.9,
        "left": 0.5
    },
    {
        "icon": "📧",
        "name": "Gmail Node",
        "desc": "Автоматизація пошти: відправка, отримання, пошук, фільтрація, вкладення",
        "color": (239, 68, 68),
        "top": 4.9,
        "left": 5.5
    }
]

for node in nodes:
    # Node box
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(node["left"]), Inches(node["top"]), Inches(4.3), Inches(1.4))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box.line.color.rgb = RGBColor(*node["color"])
    box.line.width = Pt(2)

    # Icon and name
    header = slide.shapes.add_textbox(Inches(node["left"] + 0.2), Inches(node["top"] + 0.15), Inches(4), Inches(0.4))
    tf = header.text_frame
    tf.text = f"{node['icon']} {node['name']}"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(*node["color"])

    # Description
    desc = slide.shapes.add_textbox(Inches(node["left"] + 0.2), Inches(node["top"] + 0.6), Inches(3.9), Inches(0.7))
    tf = desc.text_frame
    tf.text = node["desc"]
    tf.word_wrap = True
    tf.paragraphs[0].font.size = Pt(12)
    tf.paragraphs[0].font.color.rgb = RGBColor(75, 85, 99)

# Slide: Ви тепер теж можете кодувати
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(250, 245, 255)

title = slide.shapes.add_textbox(Inches(0.5), Inches(0.8), Inches(9), Inches(1.2))
tf = title.text_frame
tf.text = "💡 Ви тепер теж можете 'кодувати'"
tf.paragraphs[0].font.size = Pt(44)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(126, 34, 206)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Points
points = [
    "✓ Вам не потрібно бути програмістом, щоб бути 'будівельником'",
    "✓ Ви можете з'єднувати інструменти, якими вже користуєтесь",
    "✓ Ви розумієте 'мову' спілкування сервісів (API та JSON)",
    "✓ Ви розумієте різницю між ОДНИМ елементом та СПИСКОМ!"
]

y_position = 2.5
for point in points:
    point_box = slide.shapes.add_textbox(Inches(1.5), Inches(y_position), Inches(7), Inches(0.7))
    tf = point_box.text_frame
    tf.text = point
    tf.word_wrap = True
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.color.rgb = RGBColor(55, 65, 81)
    y_position += 0.9

# Question
question = slide.shapes.add_textbox(Inches(0.5), Inches(5.8), Inches(9), Inches(0.8))
tf = question.text_frame
tf.text = "Що ви автоматизуєте в першу чергу?"
tf.paragraphs[0].font.size = Pt(28)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(31, 41, 55)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Brain icon
brain = slide.shapes.add_textbox(Inches(4.5), Inches(6.5), Inches(1), Inches(0.6))
tf = brain.text_frame
tf.text = "🧠"
tf.paragraphs[0].font.size = Pt(48)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Slide: YouTube Resources with QR codes
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(239, 246, 255)

title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
tf = title.text_frame
tf.text = "📺 Корисні YouTube канали про n8n"
tf.paragraphs[0].font.size = Pt(38)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(239, 68, 68)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

subtitle = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(0.5))
stf = subtitle.text_frame
stf.text = "Відскануйте QR-код щоб підписатись"
stf.paragraphs[0].font.size = Pt(20)
stf.paragraphs[0].font.color.rgb = RGBColor(107, 114, 128)
stf.paragraphs[0].alignment = PP_ALIGN.CENTER

# QR code channels
channels = [
    {
        "name": "Roger Roger AI",
        "desc": "Практичні приклади автоматизації",
        "url": "https://youtube.com/@rogerroger-ai?si=6Es3xiXmAoJxdZwQ",
        "left": 1.2,
        "top": 2.5
    },
    {
        "name": "Kodarik",
        "desc": "Туторіали та кейси",
        "url": "https://youtube.com/@kodarik?si=nfUbXaaP09DQAuok",
        "left": 5.8,
        "top": 2.5
    },
    {
        "name": "Nate Herk",
        "desc": "Глибокі dive в n8n",
        "url": "https://youtube.com/@nateherk?si=s--1hVfNn-S3ZMvd",
        "left": 1.2,
        "top": 4.8
    },
    {
        "name": "Creator Magic AI",
        "desc": "AI + автоматизація",
        "url": "https://youtube.com/@creatormagicai?si=_WDcQcVUoYVJ4xy0",
        "left": 5.8,
        "top": 4.8
    }
]

for channel in channels:
    # Channel box
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(channel["left"]), Inches(channel["top"]), Inches(3.5), Inches(1.8))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box.line.color.rgb = RGBColor(239, 68, 68)
    box.line.width = Pt(2)

    # QR code placeholder (using QR API)
    try:
        qr_url = f"https://api.qrserver.com/v1/create-qr-code/?size=150x150&data={channel['url']}"
        response = requests.get(qr_url, timeout=5)
        if response.status_code == 200:
            image_stream = BytesIO(response.content)
            pic = slide.shapes.add_picture(image_stream, Inches(channel["left"] + 0.3), Inches(channel["top"] + 0.2), width=Inches(1.2))
    except:
        # If QR generation fails, add placeholder text
        qr_placeholder = slide.shapes.add_textbox(Inches(channel["left"] + 0.3), Inches(channel["top"] + 0.2), Inches(1.2), Inches(1.2))
        tf = qr_placeholder.text_frame
        tf.text = "QR"
        tf.paragraphs[0].font.size = Pt(24)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Channel name
    name_box = slide.shapes.add_textbox(Inches(channel["left"] + 1.6), Inches(channel["top"] + 0.3), Inches(1.7), Inches(0.5))
    tf = name_box.text_frame
    tf.text = channel["name"]
    tf.word_wrap = True
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(31, 41, 55)

    # Description
    desc_box = slide.shapes.add_textbox(Inches(channel["left"] + 1.6), Inches(channel["top"] + 0.9), Inches(1.7), Inches(0.6))
    tf = desc_box.text_frame
    tf.text = channel["desc"]
    tf.word_wrap = True
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = RGBColor(107, 114, 128)

# Final Slide: Дякую!
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(30, 27, 75)

title = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
tf = title.text_frame
tf.text = "Дякую!"
tf.paragraphs[0].font.size = Pt(72)
tf.paragraphs[0].font.bold = True
tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

subtitle = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(0.8))
stf = subtitle.text_frame
stf.text = "Запитання?"
stf.paragraphs[0].font.size = Pt(32)
stf.paragraphs[0].font.color.rgb = RGBColor(200, 200, 200)
stf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Icons
icons = slide.shapes.add_textbox(Inches(3.5), Inches(5.5), Inches(3), Inches(0.8))
tf = icons.text_frame
tf.text = "💬  ❓  💡"
tf.paragraphs[0].font.size = Pt(40)
tf.paragraphs[0].alignment = PP_ALIGN.CENTER

print("Додано фінальні слайди: ноди n8n, детальна інформація, QR-коди, дякую")

# Save
prs.save('/home/user/demo2/n8n-presentation.pptx')
print(f"✅ Презентація готова! Всього слайдів: {len(prs.slides)}")
print("📁 Файл: /home/user/demo2/n8n-presentation.pptx")
